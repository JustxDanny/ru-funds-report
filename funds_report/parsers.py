"""Pure parsers: bytes/str → list[(date, Decimal NAV)]. No I/O, no network."""

from __future__ import annotations

import io
import json
import re
from datetime import date, datetime
from decimal import Decimal

from openpyxl import load_workbook

# ── first-am.ru: full history embedded as `var chartData = [...]` JSON literal ─
FIRSTAM_CHART = re.compile(r"var\s+chartData\s*=\s*(\[.*?\]);", re.S)

# ── wealthim.ru: form pointing to authoritative xlsx export ─
WEALTHIM_FORM = re.compile(r'<form[^>]*id="export_xlsx"[^>]*>(.*?)</form>', re.S)
WEALTHIM_FIELD = re.compile(r'name="([A-Z_0-9]+)"\s+value="([^"]*)"')
DATE_DD_MM_YYYY = re.compile(r"^\d{2}\.\d{2}\.\d{4}$")


def _to_date(s: str) -> date:
    return datetime.strptime(s, "%d.%m.%Y").date()


def parse_firstam_chart(html: str) -> list[tuple[date, Decimal]]:
    """Extract the chartData array from a first-am.ru fund page.

    First-am embeds full NAV history as JSON in the page; we read the literal
    rather than the rendered table. Result is float-precision RUB → we convert via
    `str()` to preserve Decimal exactness.
    """
    m = FIRSTAM_CHART.search(html)
    if not m:
        raise ValueError("chartData not found in first-am page")
    data = json.loads(m.group(1))
    rows: list[tuple[date, Decimal]] = []
    seen: set[date] = set()
    for row in data:
        d = _to_date(row["dateFormat"])
        if d in seen:
            continue
        seen.add(d)
        rows.append((d, Decimal(str(row["price"]))))
    rows.sort(key=lambda x: x[0])
    return rows


def extract_wealthim_export_params(html: str) -> dict[str, str]:
    """Parse the hidden `<form id="export_xlsx">` and return its fields.

    These (FOND_ID, IBLOCK_ID, etc.) are the input to the xlsx export endpoint.
    """
    m = WEALTHIM_FORM.search(html)
    if not m:
        raise ValueError("export_xlsx form not found")
    fields = dict(WEALTHIM_FIELD.findall(m.group(1)))
    if "FOND_ID" not in fields:
        raise ValueError("FOND_ID missing from export form")
    return fields


def parse_wealthim_xlsx(content: bytes) -> list[tuple[date, Decimal]]:
    """Parse the wealthim NAV xlsx export. Authoritative — 4-decimal precision."""
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows: list[tuple[date, Decimal]] = []
    seen: set[date] = set()
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None or row[1] is None:
            continue
        d_str = str(row[0]).strip()
        if not DATE_DD_MM_YYYY.match(d_str):
            continue
        try:
            d = _to_date(d_str)
        except ValueError:
            continue
        try:
            v = Decimal(str(row[1]))
        except Exception:
            continue
        if d in seen:
            continue
        seen.add(d)
        rows.append((d, v))
    rows.sort(key=lambda x: x[0])
    return rows
